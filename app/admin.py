"""Administrator-only views: users, reports, audit log and backups."""
from __future__ import annotations

import json
import mimetypes
from io import BytesIO
import os
from datetime import date, datetime, timezone
from functools import wraps
from pathlib import Path

from flask import Blueprint, abort, current_app, flash, redirect, render_template, request, send_file, session, url_for
from flask_login import current_user, login_required
from werkzeug.security import generate_password_hash
from werkzeug.utils import secure_filename
from sqlalchemy import func, or_

from .extensions import db
from .services.backup import create_backup, next_backup_datetime, restore_backup
from .server_config import (
    configure_windows_firewall,
    configure_windows_service_autostart,
    schedule_windows_service_restart,
    validate_port,
)
from .utils import format_amount, gregorian_to_jalali, gregorian_datetime_to_jalali, jalali_to_gregorian_iso

admin_bp = Blueprint("admin", __name__, url_prefix="/admin")


_AUDIT_FIELD_LABELS = {
    "id": "شناسه", "entry_date": "تاریخ", "entry_type": "نوع",
    "description": "شرح", "amount": "مبلغ", "created_by_user_id": "کاربر",
    "first_name": "نام", "last_name": "نام خانوادگی", "username": "نام کاربری",
    "role": "نقش", "is_active": "وضعیت", "backup_path": "مسیر پشتیبان",
    "auto_backup_enabled": "پشتیبان‌گیری خودکار", "server_port": "پورت سرور",
    "windows_autostart_enabled": "اجرای خودکار پس از روشن شدن ویندوز",
    "path": "مسیر", "filename": "نام فایل", "previous_logo": "لوگوی قبلی", "backup_time": "ساعت بکاپ",
    "print_paper_size": "اندازه کاغذ", "print_orientation": "جهت چاپ", "success": "نتیجه",
    "server_autostart": "اجرای خودکار سرور",
}


def _audit_value_fa(key, value):
    if key == "amount":
        return format_amount(value)
    if key == "entry_type":
        return {"income": "درآمد", "expense": "هزینه"}.get(str(value), str(value))
    if key in {"is_active", "auto_backup_enabled"}:
        return "فعال" if bool(value) else "غیرفعال"
    if key == "role":
        return {"admin": "مدیر", "user": "کاربر"}.get(str(value).lower(), str(value))
    if key == "print_orientation":
        return {"landscape": "افقی", "portrait": "عمودی"}.get(str(value), str(value))
    if key == "success":
        return "موفق" if bool(value) else "ناموفق"
    if value is None:
        return "—"
    return str(value)


def _audit_data_fa(raw):
    """Turn stored technical JSON snapshots into short Persian display text."""
    if not raw:
        return ""
    try:
        data = json.loads(raw) if isinstance(raw, str) else raw
    except (TypeError, json.JSONDecodeError):
        return str(raw)
    if not isinstance(data, dict):
        return str(data)
    parts = []
    for key, value in data.items():
        # Internal ownership IDs are useful in storage but noisy in the UI.
        if key == "created_by_user_id":
            continue
        label = _AUDIT_FIELD_LABELS.get(key, key.replace("_", " "))
        parts.append(f"{label}: {_audit_value_fa(key, value)}")
    return "، ".join(parts)


def _audit_details_fa(log):
    before = _audit_data_fa(log.before_data)
    after = _audit_data_fa(log.after_data)
    message = (log.message or "").strip()
    chunks = []
    if message:
        chunks.append(message)
    if before and after:
        chunks.append(f"قبل: {before} ← بعد: {after}")
    elif before:
        chunks.append(f"اطلاعات قبلی: {before}")
    elif after:
        chunks.append(f"اطلاعات: {after}")
    return " | ".join(chunks) or "—"


def _event_label_fa(event_type):
    labels = {
        "admin_setup": "ایجاد مدیر اصلی", "login_success": "ورود کاربر", "login_failed": "ورود ناموفق",
        "login_failed_inactive": "ورود ناموفق - حساب غیرفعال", "login_failed_locked": "ورود ناموفق - حساب قفل‌شده",
        "account_locked": "قفل شدن حساب کاربر", "logout": "خروج کاربر", "create_entry": "ثبت تراکنش",
        "update_entry": "ویرایش تراکنش", "entry_date_changed": "تغییر تاریخ تراکنش",
        "entry_type_changed": "تغییر نوع تراکنش", "entry_description_changed": "تغییر شرح تراکنش",
        "entry_amount_changed": "تغییر مبلغ تراکنش", "soft_delete_entry": "حذف تراکنش",
        "user_created": "ایجاد کاربر", "user_edit": "ویرایش کاربر", "user_activate": "فعال‌سازی کاربر",
        "user_deactivate": "غیرفعال‌سازی کاربر", "user_toggle_status": "تغییر وضعیت کاربر", "user_unlock": "بازکردن قفل کاربر",
        "user_reset_password": "بازنشانی رمز عبور کاربر", "user_delete": "حذف کاربر",
        "logs_cleared": "پاک‌سازی لاگ‌ها", "app_logo_changed": "تغییر لوگوی برنامه", "app_logo_deleted": "حذف لوگوی برنامه",
        "admin_password_recovered": "بازیابی رمز مدیر اصلی", "backup_health_checked": "بررسی سلامت بکاپ",
        "backup_settings_changed": "تغییر تنظیمات پشتیبان‌گیری",
        "backup_manual": "پشتیبان‌گیری دستی", "backup_fallback": "پشتیبان‌گیری در مسیر جایگزین",
        "backup_failed": "خطا در پشتیبان‌گیری", "backup_auto": "پشتیبان‌گیری خودکار",
        "backup_manual": "پشتیبان‌گیری دستی", "backup_fallback": "پشتیبان‌گیری در مسیر جایگزین", "backup_restored": "بازیابی پشتیبان",
        "server_port_changed": "تغییر پورت سرور", "print_preferences_changed": "تغییر تنظیمات چاپ",
        "server_settings_changed": "تغییر تنظیمات سرور", "server_settings_change_failed": "خطا در تغییر تنظیمات سرور",
        "server_port_change_failed": "خطا در تغییر پورت سرور",
    }
    return labels.get(str(event_type), str(event_type))

def _related_record_fa(log):
    labels = {"DailyEntry": "ثبت روزانه", "User": "کاربر", "Setting": "تنظیمات"}
    label = labels.get(log.related_record_type, "")
    if log.related_record_id:
        return f"{label or 'رکورد'} شماره {log.related_record_id}"
    return label or "—"


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        if getattr(current_user, "role", "") not in ("admin", "Admin", "ADMIN"):
            flash("دسترسی فقط برای ادمین مجاز است.", "error")
            return redirect(url_for("entries.dashboard"))
        return view(*args, **kwargs)
    return wrapped


def _model(name):
    from . import models
    return getattr(models, name)


def _snapshot(user):
    return {"first_name": getattr(user, "first_name", ""), "last_name": getattr(user, "last_name", ""),
            "username": getattr(user, "username", "")}


def _log(event, message, success=True, related_id=None, before=None, after=None):
    try:
        AuditLog = _model("AuditLog")
        actor = current_user if getattr(current_user, "is_authenticated", False) else None
        snap = _snapshot(actor) if actor else {}
        data = dict(event_type=event, message=message, success=success,
                    actor_user_id=getattr(actor, "id", None),
                    actor_name_snapshot=(f"{snap.get('first_name', '')} {snap.get('last_name', '')}").strip(),
                    actor_username_snapshot=snap.get("username"), related_record_id=related_id,
                    before_data=json.dumps(before, ensure_ascii=False, default=str) if before else None,
                    after_data=json.dumps(after, ensure_ascii=False, default=str) if after else None,
                    client_ip=request.remote_addr, created_at=datetime.now(timezone.utc))
        db.session.add(AuditLog(**{k: v for k, v in data.items() if hasattr(AuditLog, k)}))
    except Exception:
        current_app.logger.exception("Unable to write audit log")


@admin_bp.get("/")
@admin_required
def dashboard():
    return redirect(url_for("admin.report"))


@admin_bp.get("/logo")
@admin_required
def app_logo():
    Setting = _model("Setting")
    setting = Setting.query.filter_by(key="app_logo_file").first()
    filename = str(setting.value or "").strip() if setting else ""
    if not filename:
        abort(404)
    path = Path(current_app.config["DATA_DIR"]) / "branding" / filename
    if not path.is_file():
        abort(404)
    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return send_file(path, mimetype=mime, max_age=300)

@admin_bp.route("/settings", methods=["GET", "POST"])
@admin_required
def settings():
    Setting = _model("Setting")
    port_setting = Setting.query.filter_by(key="server_port").first()
    if port_setting is None:
        port_setting = Setting(key="server_port", value=str(current_app.config.get("DEFAULT_PORT", 4000)))
        db.session.add(port_setting)
    autostart_setting = Setting.query.filter_by(key="windows_autostart_enabled").first()
    if autostart_setting is None:
        autostart_setting = Setting(key="windows_autostart_enabled", value="1")
        db.session.add(autostart_setting)
    print_paper_setting = Setting.query.filter_by(key="print_paper_size").first()
    if print_paper_setting is None:
        print_paper_setting = Setting(key="print_paper_size", value="A4")
        db.session.add(print_paper_setting)
    print_orientation_setting = Setting.query.filter_by(key="print_orientation").first()
    if print_orientation_setting is None:
        print_orientation_setting = Setting(key="print_orientation", value="landscape")
        db.session.add(print_orientation_setting)
    db.session.commit()
    logo_setting = Setting.query.filter_by(key="app_logo_file").first()
    if logo_setting is None:
        logo_setting = Setting(key="app_logo_file")
        db.session.add(logo_setting)
    db.session.commit()

    if request.method == "POST":
        if request.form.get("delete_logo") == "1":
            previous_logo = str(logo_setting.value or "").strip()
            if previous_logo:
                logo_path = Path(current_app.config["DATA_DIR"]) / "branding" / previous_logo
                try:
                    logo_path.unlink(missing_ok=True)
                except OSError:
                    pass
            logo_setting.value = ""
            db.session.commit()
            _log("app_logo_deleted", "\u0644\u0648\u06af\u0648\u06cc \u0628\u0631\u0646\u0627\u0645\u0647 \u062d\u0630\u0641 \u0634\u062f", before={"filename": previous_logo})
            flash("\u0644\u0648\u06af\u0648 \u0628\u0627 \u0645\u0648\u0641\u0642\u06cc\u062a \u062d\u0630\u0641 \u0634\u062f.", "success")
            return redirect(url_for("admin.settings"))
        uploaded_logo = request.files.get("logo_file")
        if uploaded_logo and uploaded_logo.filename:
            original_name = secure_filename(uploaded_logo.filename)
            extension = Path(original_name).suffix.lower()
            allowed = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
            if extension not in allowed:
                flash("فرمت لوگو باید PNG، JPG، JPEG یا WebP باشد.", "error")
                return redirect(url_for("admin.settings"))
            uploaded_logo.stream.seek(0, 2)
            file_size = uploaded_logo.stream.tell()
            uploaded_logo.stream.seek(0)
            if file_size > 2 * 1024 * 1024:
                flash("حجم لوگو نباید بیشتر از ۲ مگابایت باشد.", "error")
                return redirect(url_for("admin.settings"))
            branding_dir = Path(current_app.config["DATA_DIR"]) / "branding"
            branding_dir.mkdir(parents=True, exist_ok=True)
            filename = f"dailybook-logo{extension}"
            uploaded_logo.save(branding_dir / filename)
            logo_setting.value = filename
            db.session.commit()
            _log("app_logo_changed", "لوگوی برنامه تغییر کرد", after={"filename": filename})
            db.session.commit()
            flash("لوگو با موفقیت بارگذاری شد.", "success")
            return redirect(url_for("admin.settings"))
        try:
            port = validate_port(request.form.get("server_port"))
        except ValueError as exc:
            flash(str(exc), "error")
            return redirect(url_for("admin.settings"))

        autostart_enabled = request.form.get("windows_autostart_enabled") == "on"
        print_paper = request.form.get("print_paper_size", "A4")
        print_orientation = request.form.get("print_orientation", "landscape")
        if print_paper not in {"A4", "A5"} or print_orientation not in {"portrait", "landscape"}:
            flash("تنظیمات چاپ معتبر نیست.", "error")
            return redirect(url_for("admin.settings"))
        before = {
            "server_port": port_setting.value,
            "windows_autostart_enabled": autostart_setting.value == "1",
        }
        try:
            configure_windows_firewall(port)
            configure_windows_service_autostart(autostart_enabled)
        except OSError as exc:
            _log("server_settings_change_failed", str(exc), success=False, before=before,
                 after={"server_port": str(port), "windows_autostart_enabled": autostart_enabled})
            db.session.commit()
            flash(f"تنظیمات ذخیره نشد: {exc}", "error")
            return redirect(url_for("admin.settings"))

        port_changed = port_setting.value != str(port)
        port_setting.value = str(port)
        autostart_setting.value = "1" if autostart_enabled else "0"
        print_paper_setting.value = print_paper
        print_orientation_setting.value = print_orientation
        _log("server_settings_changed", "تنظیمات سرور تغییر کرد", before=before,
             after={"server_port": port_setting.value, "windows_autostart_enabled": autostart_enabled,
                    "print_paper_size": print_paper, "print_orientation": print_orientation})
        db.session.commit()

        restarted = schedule_windows_service_restart() if port_changed else False
        if restarted:
            flash(f"تنظیمات ذخیره شد؛ سرویس با پورت {port} تا چند ثانیه دیگر راه‌اندازی مجدد می‌شود.", "success")
        else:
            flash("تنظیمات ذخیره شد.", "success")
        return redirect(url_for("admin.settings"))

    return render_template(
        "admin/settings.html",
        server_port=port_setting.value,
        windows_autostart_enabled=autostart_setting.value != "0",
        environment_override=bool(os.environ.get("DAILYBOOK_PORT")),
        logo_filename=logo_setting.value,
    )


@admin_bp.route("/users", methods=["GET", "POST"])
@admin_required
def users():
    User = _model("User")
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        if not username or len(username) < 3 or len(password) < 8:
            flash("نام کاربری حداقل ۳ و رمز عبور حداقل ۸ نویسه باشد.", "error")
            return redirect(url_for("admin.users"))
        if User.query.filter_by(username=username).first():
            flash("نام کاربری تکراری است.", "error")
        else:
            user = User(first_name=request.form.get("first_name", "").strip(), last_name=request.form.get("last_name", "").strip(),
                        username=username, password_hash=generate_password_hash(password),
                        role="user", status="active")
            db.session.add(user); db.session.flush()
            _log("user_created", "کاربر ایجاد شد", related_id=user.id, after=_snapshot(user)); db.session.commit()
            flash("کاربر ایجاد شد.", "success")
        return redirect(url_for("admin.users"))
    return render_template("admin/users.html", users=User.query.order_by(User.id.desc()).all())


@admin_bp.post("/users/<int:user_id>/<action>")
@admin_required
def user_action(user_id, action):
    User = _model("User"); user = db.session.get(User, user_id)
    if not user:
        flash("کاربر پیدا نشد.", "error"); return redirect(url_for("admin.users"))
    if str(user.role).lower() == "admin" and action in {"deactivate", "delete", "toggle-status"}:
        flash("مدیر اصلی را نمی‌توان غیرفعال یا حذف کرد.", "error")
        return redirect(url_for("admin.users"))
    before = _snapshot(user) | {"status": user.status, "is_locked": user.is_locked}
    if action == "edit":
        first_name = request.form.get("first_name", "").strip()
        last_name = request.form.get("last_name", "").strip()
        username = request.form.get("username", "").strip()
        duplicate = User.query.filter(User.username == username, User.id != user.id).first()
        if not first_name or not last_name or len(username) < 3 or duplicate:
            flash("مشخصات نامعتبر است یا نام کاربری تکراری است.", "error")
            return redirect(url_for("admin.users"))
        user.first_name, user.last_name, user.username = first_name, last_name, username
    elif action == "activate": user.status = "active"
    elif action == "deactivate": user.status = "inactive"
    elif action == "toggle-status": user.status = "inactive" if user.status == "active" else "active"
    elif action == "unlock":
        user.is_locked = False; user.failed_login_count = 0
        if user.status == "locked": user.status = "active"
    elif action == "reset-password":
        password = request.form.get("password", "")
        if len(password) < 8:
            flash("رمز عبور حداقل ۸ نویسه باشد.", "error")
            return redirect(url_for("admin.users"))
        user.password_hash = generate_password_hash(password)
    elif action == "delete": user.soft_delete()
    else: flash("عملیات نامعتبر است.", "error"); return redirect(url_for("admin.users"))
    after = _snapshot(user) | {"status": user.status, "is_locked": user.is_locked}
    _log(f"user_{action.replace('-', '_')}", f"عملیات {action} روی کاربر انجام شد",
         related_id=user.id, before=before, after=after)
    db.session.commit(); flash("عملیات با موفقیت انجام شد.", "success")
    return redirect(url_for("admin.users"))



def _print_preferences():
    Setting = _model("Setting")
    paper = Setting.query.filter_by(key="print_paper_size").first()
    orientation = Setting.query.filter_by(key="print_orientation").first()
    paper_value = paper.value if paper and paper.value in {"A4", "A5"} else "A4"
    orientation_value = orientation.value if orientation and orientation.value in {"portrait", "landscape"} else "landscape"
    return paper_value, orientation_value
def _apply_report_filters(query, Entry, params):
    advanced = params.get("advanced_filter") == "on"
    state = {
        "advanced_filter": advanced,
        "use_date": params.get("use_date") == "on",
        "use_time": params.get("use_time") == "on",
        "use_user": params.get("use_user") == "on",
        "use_description": params.get("use_description") == "on",
        "start": params.get("start", "").strip(),
        "end": params.get("end", "").strip(),
        "time_filter": params.get("time_filter", "").strip(),
        "user_filter": params.get("user_filter", "").strip(),
        "description_filter": params.get("description_filter", "").strip(),
        "debit_filter": params.get("debit") == "on",
        "credit_filter": params.get("credit") == "on",
    }
    if not advanced:
        return query, state, None
    try:
        if state["use_date"]:
            start_date = date.fromisoformat(jalali_to_gregorian_iso(state["start"])) if state["start"] else None
            end_date = date.fromisoformat(jalali_to_gregorian_iso(state["end"])) if state["end"] else None
            if start_date and end_date and start_date > end_date:
                raise ValueError("تاریخ شروع نمی‌تواند بعد از تاریخ پایان باشد.")
            if start_date:
                query = query.filter(Entry.entry_date >= start_date)
            if end_date:
                query = query.filter(Entry.entry_date <= end_date)
        if state["use_time"] and state["time_filter"]:
            datetime.strptime(state["time_filter"], "%H:%M")
            query = query.filter(func.strftime("%H:%M", Entry.created_at) == state["time_filter"])
        if state["use_user"] and state["user_filter"]:
            try:
                user_id = int(state["user_filter"])
            except (TypeError, ValueError):
                raise ValueError("کاربر انتخاب‌شده معتبر نیست.") from None
            query = query.filter(Entry.created_by_user_id == user_id)
        if state["use_description"] and state["description_filter"]:
            query = query.filter(Entry.description.ilike(f"%{state['description_filter']}%"))
        types = []
        if state["debit_filter"]:
            types.append("expense")
        if state["credit_filter"]:
            types.append("income")
        if len(types) == 1:
            query = query.filter(Entry.entry_type == types[0])
    except ValueError as exc:
        return query, state, str(exc)
    return query, state, None

@admin_bp.get("/report")
@admin_required
def report():
    Entry = _model("DailyEntry")
    User = _model("User")
    users = User.query.filter(User.status != "deleted").order_by(User.first_name.asc(), User.last_name.asc()).all()
    query = Entry.query.filter(Entry.deleted_at.is_(None))
    query, filter_state, report_error = _apply_report_filters(query, Entry, request.args)
    entries = [] if report_error else query.order_by(Entry.entry_date.desc(), Entry.id.desc()).all()
    debit = sum(int(getattr(e, "amount", 0) or 0) for e in entries if getattr(e, "entry_type", "") == "expense")
    credit = sum(int(getattr(e, "amount", 0) or 0) for e in entries if getattr(e, "entry_type", "") == "income")
    return render_template("admin/report.html", entries=entries, debit=debit, credit=credit,
                           report_error=report_error, users=users, **filter_state,
                           format_amount=format_amount,
                           gregorian_to_jalali=gregorian_to_jalali,
                           gregorian_datetime_to_jalali=gregorian_datetime_to_jalali, print_paper_size=_print_preferences()[0], print_orientation=_print_preferences()[1])

@admin_bp.get("/logs")
@admin_required
def logs():
    AuditLog = _model("AuditLog")
    User = _model("User")
    advanced = request.args.get("advanced_filter") == "on"
    state = {"advanced_filter": advanced, "use_date": request.args.get("use_date") == "on", "use_time": request.args.get("use_time") == "on", "use_user": request.args.get("use_user") == "on", "use_description": request.args.get("use_description") == "on", "start": request.args.get("start", "").strip(), "end": request.args.get("end", "").strip(), "time_filter": request.args.get("time_filter", "").strip(), "user_filter": request.args.get("user_filter", "").strip(), "description_filter": request.args.get("description_filter", "").strip(), "success_filter": request.args.get("success") == "on", "failure_filter": request.args.get("failure") == "on"}
    query = AuditLog.query
    log_error = None
    if advanced:
        try:
            if state["use_date"]:
                start_date = date.fromisoformat(jalali_to_gregorian_iso(state["start"])) if state["start"] else None
                end_date = date.fromisoformat(jalali_to_gregorian_iso(state["end"])) if state["end"] else None
                if start_date and end_date and start_date > end_date: raise ValueError("تاریخ شروع نمی‌تواند بعد از تاریخ پایان باشد.")
                if start_date: query = query.filter(AuditLog.created_at >= datetime.combine(start_date, datetime.min.time(), timezone.utc))
                if end_date: query = query.filter(AuditLog.created_at <= datetime.combine(end_date, datetime.max.time(), timezone.utc))
            if state["use_time"] and state["time_filter"]:
                datetime.strptime(state["time_filter"], "%H:%M")
                query = query.filter(func.strftime("%H:%M", AuditLog.created_at) == state["time_filter"])
            if state["use_user"] and state["user_filter"]: query = query.filter(AuditLog.actor_user_id == int(state["user_filter"]))
            if state["use_description"] and state["description_filter"]:
                needle = f"%{state['description_filter']}%"
                query = query.filter(db.or_(AuditLog.event_type.ilike(needle), AuditLog.message.ilike(needle)))
            states = ([True] if state["success_filter"] else []) + ([False] if state["failure_filter"] else [])
            if len(states) == 1: query = query.filter(AuditLog.success == states[0])
        except (TypeError, ValueError): log_error = "اطلاعات فیلتر لاگ معتبر نیست."
    users = User.query.filter(User.status != "deleted").order_by(User.first_name.asc(), User.last_name.asc()).all()
    return render_template("admin/logs.html", logs=[] if log_error else query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc()).limit(500).all(), users=users, log_error=log_error, event_label_fa=_event_label_fa, **state,
                           audit_details_fa=_audit_details_fa, related_record_fa=_related_record_fa,

                           gregorian_datetime_to_jalali=gregorian_datetime_to_jalali, print_paper_size=_print_preferences()[0], print_orientation=_print_preferences()[1])


@admin_bp.post("/logs/clear")
@admin_required
def clear_logs():
    password = str(request.form.get("admin_password") or "")
    if not password or not current_user.check_password(password):
        flash("رمز مدیر نادرست است؛ لاگ‌ها پاک نشدند.", "error")
        return redirect(url_for("admin.logs"))
    AuditLog = _model("AuditLog")
    count = AuditLog.query.delete(synchronize_session=False)
    _log("logs_cleared", f"{count} لاگ پاک شد")
    db.session.commit()
    flash("لاگ‌ها پاک شدند و رویداد پاک‌سازی ثبت شد.", "success")
    return redirect(url_for("admin.logs"))


@admin_bp.post("/backup/pick-folder")
@admin_required
def pick_backup_folder():
    """Open the local Windows folder chooser for the backup destination."""
    payload = request.get_json(silent=True) or {}
    initial_path = str(payload.get("current_path") or "").strip()
    if not initial_path or not os.path.isdir(initial_path):
        initial_path = str(current_app.config.get("DEFAULT_BACKUP_DIR") or "")

    root = None
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        root.update()
        selected = filedialog.askdirectory(
            parent=root,
            initialdir=initial_path or None,
            title="انتخاب مسیر ذخیره نسخه پشتیبان DailyBook",
        )
    except Exception:
        return {"error": "باز کردن پنجره انتخاب پوشه ممکن نشد."}, 500
    finally:
        if root is not None:
            root.destroy()

    return {"path": selected or ""}

@admin_bp.route("/backup", methods=["GET", "POST"])
@admin_required
def backup():
    Setting = _model("Setting")
    setting = Setting.query.filter_by(key="backup").first()
    if setting is None:
        setting = Setting(key="backup", auto_backup_enabled=True)
        db.session.add(setting)
    time_setting = Setting.query.filter_by(key="backup_time").first()
    if time_setting is None:
        time_setting = Setting(key="backup_time", value="18:00")
        db.session.add(time_setting)
    db.session.commit()
    if request.method == "POST":
        preferred = request.form.get("backup_path", "").strip() or current_app.config.get("DEFAULT_BACKUP_DIR")
        if request.form.get("action") == "save-settings":
            backup_time = request.form.get("backup_time", "18:00").strip()
            try:
                hh, mm = (int(part) for part in backup_time.split(":", 1))
                if not (0 <= hh <= 23 and 0 <= mm <= 59): raise ValueError
            except (TypeError, ValueError):
                flash("ساعت بکاپ معتبر نیست.", "error")
                return redirect(url_for("admin.backup"))
            backup_time = f"{hh:02d}:{mm:02d}"
            before = {"backup_path": setting.backup_path, "auto_backup_enabled": setting.auto_backup_enabled, "backup_time": time_setting.value}
            setting.backup_path = str(preferred)
            setting.auto_backup_enabled = request.form.get("auto_backup_enabled") == "on"
            time_setting.value = backup_time
            after = {"backup_path": setting.backup_path, "auto_backup_enabled": setting.auto_backup_enabled, "backup_time": time_setting.value}
            _log("backup_settings_changed", "تنظیمات بکاپ تغییر کرد", before=before, after=after)
            db.session.commit()
            flash("تنظیمات بکاپ ذخیره شد.", "success")
            return redirect(url_for("admin.backup"))
        if request.form.get("action") == "check-health":
            uploaded = request.files.get("restore_file")
            if not uploaded or not uploaded.filename:
                flash("ابتدا فایل بکاپ را انتخاب کنید.", "error")
                return redirect(url_for("admin.backup"))
            extension = Path(secure_filename(uploaded.filename)).suffix.lower()
            if extension not in {".db", ".sqlite", ".sqlite3"}:
                flash("فقط فایل SQLite با پسوند DB یا SQLite قابل بررسی است.", "error")
                return redirect(url_for("admin.backup"))
            session["last_restore_path"] = secure_filename(uploaded.filename)
            health_dir = Path(current_app.config["DATA_DIR"]) / "restore"
            health_dir.mkdir(parents=True, exist_ok=True)
            health_path = health_dir / "health-check.tmp.db"
            try:
                import sqlite3
                uploaded.save(health_path)
                connection = sqlite3.connect(health_path)
                try:
                    result = connection.execute("PRAGMA integrity_check").fetchone()
                finally:
                    connection.close()
                if not result or str(result[0]).lower() != "ok":
                    raise sqlite3.DatabaseError("بررسی سلامت فایل موفق نبود.")
                flash("بکاپ سالم است و برای بازیابی آماده است.", "success")
            except Exception as exc:
                flash(f"این بکاپ سالم نیست یا قابل بررسی نیست: {exc}", "error")
            finally:
                try:
                    health_path.unlink(missing_ok=True)
                except PermissionError:
                    pass
            return redirect(url_for("admin.backup"))
        if request.form.get("action") == "restore":
            uploaded = request.files.get("restore_file")
            if not uploaded or not uploaded.filename:
                flash("فایل بکاپ را انتخاب کنید.", "error")
                return redirect(url_for("admin.backup"))
            extension = Path(secure_filename(uploaded.filename)).suffix.lower()
            if extension not in {".db", ".sqlite", ".sqlite3"}:
                flash("فقط فایل SQLite با پسوند DB یا SQLite قابل بازیابی است.", "error")
                return redirect(url_for("admin.backup"))
            session["last_restore_path"] = secure_filename(uploaded.filename)
            restore_dir = Path(current_app.config["DATA_DIR"]) / "restore"
            restore_dir.mkdir(parents=True, exist_ok=True)
            restore_path = restore_dir / "selected-backup.tmp.db"
            try:
                uploaded.save(restore_path)
                db.session.remove()
                db.engine.dispose()
                restore_backup(restore_path, current_app.config["SQLALCHEMY_DATABASE_URI"])
                _log("backup_restored", "بکاپ با موفقیت بازیابی شد", after={"path": secure_filename(uploaded.filename)})
                db.session.commit()
                flash("بکاپ با موفقیت بازیابی شد.", "success")
            except Exception as exc:
                db.session.rollback()
                flash(f"بازیابی بکاپ انجام نشد: {exc}", "error")
            finally:
                restore_path.unlink(missing_ok=True)
            return redirect(url_for("admin.backup"))
        try:
            result = create_backup(current_app.config["SQLALCHEMY_DATABASE_URI"], preferred,
                                   current_app.config["DEFAULT_BACKUP_DIR"])
            _log("backup_fallback" if result["used_fallback"] else "backup_manual", "پشتیبان ایجاد شد", after={"path": str(result["path"])})
            db.session.commit()
            if result["used_fallback"]:
                flash(f"مسیر سفارشی در دسترس نبود؛ بکاپ در مسیر پیش‌فرض ذخیره شد: {result['path']}", "error")
            else:
                flash(f"پشتیبان ایجاد شد: {result['path']}", "success")
        except Exception as exc:
            _log("backup_failed", str(exc), success=False); db.session.commit(); flash(str(exc), "error")
        return redirect(url_for("admin.backup"))
    next_backup = next_backup_datetime(setting.auto_backup_enabled, time_setting.value or "18:00", setting.value)
    AuditLog = _model("AuditLog")
    backup_events = {"backup_manual", "backup_fallback", "backup_restored", "backup_auto", "backup_auto_fallback", "backup_failed", "backup_auto_failed"}
    backup_history = []
    for event in AuditLog.query.filter(AuditLog.event_type.in_(backup_events)).order_by(
        AuditLog.created_at.desc(), AuditLog.id.desc()
    ).limit(50).all():
        after = event.after if isinstance(event.after, dict) else {}
        raw_path = str(after.get("path") or "").strip()
        file_path = Path(raw_path) if raw_path else None
        size = "—"
        if file_path and file_path.exists():
            bytes_count = file_path.stat().st_size
            size = f"{bytes_count / (1024 * 1024):.1f} مگابایت" if bytes_count >= 1024 * 1024 else f"{max(1, bytes_count // 1024)} کیلوبایت"
        event_type = str(event.event_type or "")
        if event_type in {"backup_auto", "backup_auto_fallback"}:
            kind = "خودکار"
        elif event_type in {"backup_manual", "backup_fallback"}:
            kind = "دستی"
        elif event_type == "backup_restored":
            kind = "بازیابی"
            kind = "خطا"
        backup_history.append({
            "date": gregorian_datetime_to_jalali(event.created_at),
            "kind": kind,
            "path": raw_path or (event.message or "—"),
            "filename": file_path.name if file_path else "—",
            "size": size,
            "success": bool(event.success),
            "status": "موفق" if event.success else "ناموفق",
        })
    return render_template("admin/backup.html", default_path=current_app.config.get("DEFAULT_BACKUP_DIR"),
                           setting=setting, backup_time=time_setting.value or "18:00",
                           last_backup=setting.value, next_backup=next_backup,
                           backup_history=backup_history,
                           restore_path=session.get("last_restore_path", ""))

def _safe_excel_text(value):
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in {"=", "+", "-", "@"} else text


def _excel_download(title, headers, rows, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E5FA9")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append([_safe_excel_text(value) if isinstance(value, str) else value for value in row])
    thin_side = Side(style="thin", color="B7C5D4")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 3, 14), 42)
        for cell in column[1:]:
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@admin_bp.get("/report/export.xlsx")
@admin_required
def export_report_xlsx():
    Entry = _model("DailyEntry")
    query = Entry.query.filter(Entry.deleted_at.is_(None))
    query, filter_state, filter_error = _apply_report_filters(query, Entry, request.args)
    if filter_error:
        flash(filter_error, "error")
        return redirect(url_for("admin.report", **request.args.to_dict(flat=True)))
    rows = [[gregorian_to_jalali(item.entry_date), gregorian_datetime_to_jalali(item.created_at).split(" — ")[1],
             (item.created_by_first_name_snapshot + " " + item.created_by_last_name_snapshot).strip(),
             item.description, format_amount(item.amount) if item.entry_type == "expense" else "",
             format_amount(item.amount) if item.entry_type == "income" else ""]
            for item in query.order_by(Entry.entry_date.desc(), Entry.id.desc()).all()]
    return _excel_download("گزارش دفتر", ["تاریخ", "ساعت", "کاربر", "شرح", "بدهکار", "بستانکار"],
                           rows, "DailyBook-Report.xlsx")

@admin_bp.get("/logs/export.xlsx")
@admin_required
def export_logs_xlsx():
    AuditLog = _model("AuditLog")
    query = AuditLog.query
    if request.args.get("advanced_filter") == "on":
        try:
            if request.args.get("use_date") == "on":
                start = request.args.get("start", "").strip()
                end = request.args.get("end", "").strip()
                start_date = date.fromisoformat(jalali_to_gregorian_iso(start)) if start else None
                end_date = date.fromisoformat(jalali_to_gregorian_iso(end)) if end else None
                if start_date and end_date and start_date > end_date: raise ValueError("تاریخ شروع نمی‌تواند بعد از تاریخ پایان باشد.")
                if start_date: query = query.filter(AuditLog.created_at >= datetime.combine(start_date, datetime.min.time(), timezone.utc))
                if end_date: query = query.filter(AuditLog.created_at <= datetime.combine(end_date, datetime.max.time(), timezone.utc))
            if request.args.get("use_time") == "on" and request.args.get("time_filter"):
                datetime.strptime(request.args["time_filter"], "%H:%M")
                query = query.filter(func.strftime("%H:%M", AuditLog.created_at) == request.args["time_filter"])
            if request.args.get("use_user") == "on" and request.args.get("user_filter"):
                query = query.filter(AuditLog.actor_user_id == int(request.args["user_filter"]))
            if request.args.get("use_description") == "on" and request.args.get("description_filter"):
                needle = f"%{request.args['description_filter']}%"
                query = query.filter(db.or_(AuditLog.event_type.ilike(needle), AuditLog.message.ilike(needle)))
            states = ([True] if request.args.get("success") == "on" else []) + ([False] if request.args.get("failure") == "on" else [])
            if len(states) == 1: query = query.filter(AuditLog.success == states[0])
        except (TypeError, ValueError):
            flash("اطلاعات فیلتر لاگ معتبر نیست.", "error")
            return redirect(url_for("admin.logs", **request.args.to_dict(flat=True)))
    else:
        if request.args.get("event"): query = query.filter_by(event_type=request.args["event"])
        if request.args.get("user_id"): query = query.filter_by(actor_user_id=request.args["user_id"])
    rows = [[gregorian_datetime_to_jalali(item.created_at), _event_label_fa(item.event_type), item.actor_name_snapshot or "", item.client_ip or "", _related_record_fa(item), "موفق" if item.success else "ناموفق", _audit_details_fa(item)] for item in query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc()).limit(500).all()]
    return _excel_download("لاگ فعالیت‌ها", ["زمان", "رویداد", "کاربر", "IP سیستم", "مورد مرتبط", "نتیجه", "جزئیات"], rows, "DailyBook-Activity-Logs.xlsx")
